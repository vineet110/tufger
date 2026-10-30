"""Turn an uploaded Excel / CSV / vCard file into a clean player list."""

import csv
import io
import quopri
import re
import unicodedata
from itertools import islice

from openpyxl import load_workbook

MAX_ROWS = 5000  # also caps how far a zip-bomb .xlsx gets read

NAME_HINTS = {"name", "names", "player", "players", "naam"}
PHONE_HINTS = {"phone", "mobile", "mob", "number", "contact", "whatsapp", "cell", "ph"}  # not a bare "no": "Sr No", "Jersey No"
NOT_PHONE = {"label", "type", "sr", "sl", "serial", "sno"}  # "Phone 1 - Label", "Sr. No."
PHONE_IN_TEXT = re.compile(r"(?:\+?91[\s-]?|0)?\d(?:[\s-]?\d){9}")


class RosterError(ValueError):
    """A problem with the file itself; the message is safe to show the host."""


def normalize_phone(raw: object) -> str | None:
    """The 10 digits of a number, for matching. New numbers also go through mobile_phone."""
    if raw is None:
        return None
    if isinstance(raw, float):
        if not raw.is_integer():
            return None
        raw = str(int(raw))  # Excel stores 9876543210 as 9876543210.0
    digits = re.sub(r"\D", "", str(raw))
    if len(digits) == 14 and digits.startswith("0091"):
        digits = digits[4:]
    elif len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    return digits if len(digits) == 10 else None


def mobile_phone(raw: object) -> str | None:
    """A number that can be saved: older rows may hold 10 digits that wouldn't pass this."""
    phone = normalize_phone(raw)
    return phone if phone and phone[0] in "6789" else None


def _strip_invisible(text: str) -> str:
    """Zero-width characters and emoji leftovers (variation selectors, keycaps) from WhatsApp."""
    return "".join(c for c in text if unicodedata.category(c) != "Cf" and c != "\u20e3" and not "\ufe00" <= c <= "\ufe0f")


def _name_char(c: str) -> str:
    return c if unicodedata.category(c)[0] in "LMN" or c in " .'-" else " "


def clean_name(raw: object) -> str:
    """Tidy a name. Short numbers stay ("Rahul 2") because they tell two people apart; phone numbers don't."""
    text = _strip_invisible(unicodedata.normalize("NFKC", str(raw or "")))
    text = PHONE_IN_TEXT.sub(" ", text)
    text = re.sub(r"\d{5,}", " ", text)
    text = re.sub(r"^\s*\d+[\s.)\-:]+", " ", text)  # a serial number in front
    text = " ".join("".join(_name_char(c) for c in text).split()).strip(" .-'")
    return text.title()[:120] if any(c.isalpha() for c in text) else ""


def name_key(name: str) -> str:
    """What makes two names "the same": case, spacing and punctuation don't count."""
    text = _strip_invisible(unicodedata.normalize("NFKC", str(name))).casefold()
    words = "".join(c if unicodedata.category(c)[0] in "LMN" else " " for c in text).split()
    return " ".join(words) or " ".join(text.split())


def _decode(content: bytes) -> str:
    if content[:2] in (b"\xff\xfe", b"\xfe\xff"):  # Excel's "Unicode text" export
        return content.decode("utf-16", errors="ignore")
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="ignore")


def _read_rows(filename: str, content: bytes) -> list[tuple[int, list[object]]]:
    """Non-blank rows with their row number in the file, so problems can be pointed at."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = wb.worksheets[0].iter_rows(values_only=True)
    else:
        text = _decode(content).replace("\x00", "")
        first = next((line for line in text.splitlines() if line.strip()), "")
        delimiter = max(",;\t", key=first.count)
        rows = csv.reader(io.StringIO(text), delimiter=delimiter)
    numbered = list(islice(enumerate(rows, start=1), MAX_ROWS + 1))
    if len(numbered) > MAX_ROWS:
        raise RosterError(f"That file has more than {MAX_ROWS} rows. Split it and upload the parts")
    return [(n, list(row)) for n, row in numbered if any(c not in (None, "") for c in row)]


def _words(cell: object) -> set[str]:
    label = re.sub(r"([a-z])([A-Z])", r"\1 \2", str(cell or "")).lower()  # FirstName
    label = re.sub(r"(first|last|full|given|family|player|contact|sur)name", r"\1 name", label)
    label = re.sub(r"(phone|mobile|contact|whatsapp|cell)(number|num|no)\b", r"\1 \2", label)
    return set(re.findall(r"[a-z]+", label))


def _header_columns(row: list[object]) -> tuple[list[int], int | None]:
    """Read a heading row: (name columns, phone column)."""
    labels = [_words(c) for c in row]
    if any(normalize_phone(c) for c in row):
        return [], None  # a row with a number in it is data
    firsts = [i for i, w in enumerate(labels) if "name" in w and w & {"first", "given"} and "phonetic" not in w]
    lasts = [i for i, w in enumerate(labels) if "name" in w and w & {"last", "family", "sur"} and "phonetic" not in w]
    if firsts:
        name_cols = firsts[:1] + lasts[:1]  # Google and Outlook split the name
    else:
        named = [i for i, w in enumerate(labels) if w & NAME_HINTS]
        name_cols = [i for i in named if not labels[i] & (PHONE_HINTS - {"contact"})][:1] or named[:1]
    phones = [
        i for i, w in enumerate(labels)
        if w & PHONE_HINTS and not w & NOT_PHONE and w != {"s", "no"} and i not in name_cols
    ]
    phones.sort(key=lambda i: not labels[i] & {"mobile", "whatsapp", "cell"})  # sort is stable: mobile first
    return name_cols, phones[0] if phones else None


def _find_columns(rows: list[list[object]]) -> tuple[int, list[int], int | None]:
    """Return (first_data_row, name_cols, phone_col)."""
    found = None
    for idx, row in enumerate(rows[:5]):
        name_cols, phone_col = _header_columns(row)
        if name_cols and phone_col is not None:
            return idx + 1, name_cols, phone_col
        if found is None and name_cols:
            found = (idx + 1, name_cols, phone_col)
        elif found is None and phone_col is not None:
            # only a number heading: "Sid (no number)" is a person, so the rows below must hold numbers
            below = rows[idx + 1:idx + 6]
            if any(phone_col < len(r) and normalize_phone(r[phone_col]) for r in below):
                found = (idx + 1, name_cols, phone_col)
    if found:
        return found

    # No header row: guess from the data itself.
    width = max((len(r) for r in rows), default=0)
    phone_hits = [0] * width
    text_hits = [0] * width
    for row in rows[:50]:
        for i, cell in enumerate(row):
            if normalize_phone(cell):
                phone_hits[i] += 1
            elif clean_name(cell):
                text_hits[i] += 1
    phone_col = max(range(width), key=lambda i: phone_hits[i]) if width and max(phone_hits) else None
    name_col = max(range(width), key=lambda i: text_hits[i]) if width and max(text_hits) else None
    if name_col is None:
        name_col = phone_col  # one column of "Rahul 98765 43210"
    return 0, [] if name_col is None else [name_col], phone_col


def _read_vcard(content: bytes) -> list[dict]:
    """Phone contact exports. A card with several numbers uses the mobile one."""
    text = re.sub(r"\r?\n[ \t]", "", _decode(content))
    lines: list[str] = []
    for line in text.splitlines():  # quoted-printable values wrap with a trailing "="
        if lines and lines[-1].endswith("=") and "QUOTED-PRINTABLE" in lines[-1].partition(":")[0].upper():
            lines[-1] = lines[-1][:-1] + line
        else:
            lines.append(line)
    text = "\n".join(lines)
    contacts: list[dict] = []
    for card_no, card in enumerate(re.findall(r"BEGIN:VCARD(.*?)END:VCARD", text, re.S | re.I), start=1):
        name, fallback, phones = "", "", []
        for line in card.splitlines():
            head, _, value = line.partition(":")
            params = head.upper().split(";")
            field = params[0].rsplit(".", 1)[-1]
            if "ENCODING=QUOTED-PRINTABLE" in params:
                value = quopri.decodestring(value.encode()).decode("utf-8", errors="ignore")
            value = value.replace("\\,", ",").replace("\\;", ";").strip()
            if field == "FN":
                name = value
            elif field == "N":
                fallback = " ".join(reversed([p for p in value.split(";")[:2] if p]))
            elif field == "TEL" and value:
                mobile = "CELL" in head.upper() or "MOBILE" in head.upper()
                phones.append((not mobile_phone(value), not mobile, value))
        phones.sort(key=lambda p: p[:2])  # a good number first, then the mobile one
        contacts.append({"line": card_no, "name": name or fallback, "phone": phones[0][2] if phones else None})
    if len(contacts) > MAX_ROWS:
        raise RosterError(f"That file has more than {MAX_ROWS} contacts. Split it and upload the parts")
    return contacts


def read_contacts(filename: str, content: bytes) -> list[dict]:
    """Every row as it stands in the file: {line, name, phone}, nothing cleaned or removed yet."""
    if filename.lower().endswith(".vcf"):
        return _read_vcard(content)
    numbered = _read_rows(filename, content)
    if not numbered:
        return []
    start, name_cols, phone_col = _find_columns([row for _, row in numbered])

    contacts: list[dict] = []
    for line, row in numbered[start:]:
        cell = lambda i: row[i] if i is not None and i < len(row) else None  # noqa: E731
        name = " ".join(str(cell(i)) for i in name_cols if cell(i) not in (None, ""))
        contacts.append({"line": line, "name": name, "phone": cell(phone_col)})
    return contacts


def parse_sheet(filename: str, content: bytes) -> list[dict]:
    """A one-off squad for a booking: unreadable numbers are dropped and repeats removed."""
    players: list[dict] = []
    seen: set[str] = set()
    for row in read_contacts(filename, content):
        phone = normalize_phone(row["phone"])
        name = clean_name(row["name"])
        if not name and not phone:
            continue
        key = phone or name_key(name)
        if key in seen:
            continue
        seen.add(key)
        players.append({"name": name or f"Player {phone[-4:]}", "phone": phone})
    return players
